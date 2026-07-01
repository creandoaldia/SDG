"""
SDG Localidades — Generador de Excel individual por tab
Produce un Excel ligero con los datos procesados de un solo tipo de fuente.
Sin portada, sin validacion, sin secciones de informe completo.
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


class TabExcelWriter:
    """
    Escribe un Excel individual para un tipo de tab (PQRS, SAC, CR, etc.).
    Hoja 1: datos procesados/normalizados
    Hoja 2 (opcional): resumen o tabla adicional
    """

    C = {
        'header': '1F4E79',
        'subheader': '2E75B6',
        'accent': 'D6E4F0',
    }

    font_hdr = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    font_data = Font(name='Calibri', size=10)
    font_title = Font(name='Calibri', size=12, bold=True, color='1F4E79')
    fill_hdr = PatternFill(start_color=C['header'], end_color=C['header'], fill_type='solid')
    fill_sub = PatternFill(start_color=C['subheader'], end_color=C['subheader'], fill_type='solid')
    fill_accent = PatternFill(start_color=C['accent'], end_color=C['accent'], fill_type='solid')

    def write(self, tab_type: str, ingestion_result, normalized: dict,
              extra_data: dict, output_path: str, month: str, year: str):
        """Genera el Excel individual del tab."""
        wb = Workbook()

        # ── Hoja 1: Datos procesados ──
        ws_data = wb.active
        ws_data.title = "DATOS PROCESADOS"

        # Titulo
        titles = {
            "pqrs": "PQRS", "atenciones": "ATENCIONES SAC",
            "cert-residencia": "CERTIFICADO RESIDENCIA",
            "prop-horizontal": "PROPIEDAD HORIZONTAL",
            "encuestas": "ENCUESTAS",
            "doc-extraviados": "DOCUMENTOS EXTRAVIADOS"
        }
        title = titles.get(tab_type, tab_type.upper())
        ws_data.merge_cells('A1:Z1')
        ws_data['A1'] = f"{title} — {month} {year}"
        ws_data['A1'].font = self.font_title
        ws_data['A1'].alignment = Alignment(horizontal='left')

        # Escribir cada hoja normalizada
        row_offset = 3
        for sheet_name, df in normalized.items():
            if df is None or df.empty:
                continue

            # Subheader con nombre de hoja
            ws_data.merge_cells(start_row=row_offset, start_column=1,
                                end_row=row_offset, end_column=min(len(df.columns), 10))
            ws_data.cell(row=row_offset, column=1, value=f"Hoja: {sheet_name}")
            ws_data.cell(row=row_offset, column=1).font = Font(name='Calibri', size=10, bold=True, color='2E75B6')
            row_offset += 1

            # Headers
            for col_idx, col_name in enumerate(df.columns, 1):
                cell = ws_data.cell(row=row_offset, column=col_idx, value=str(col_name)[:60])
                cell.font = self.font_hdr
                cell.fill = self.fill_hdr
                cell.alignment = Alignment(horizontal='center')
            row_offset += 1

            # Datos
            for _, row_data in df.iterrows():
                for col_idx, val in enumerate(row_data, 1):
                    cell = ws_data.cell(row=row_offset, column=col_idx)
                    if pd.notna(val):
                        cell.value = str(val)[:255]
                    cell.font = self.font_data
                row_offset += 1

            row_offset += 1  # espacio entre hojas

        # ── Hoja 2: Resumen (si hay datos extra) ──
        if extra_data:
            ws_sum = wb.create_sheet("RESUMEN")
            ws_sum.merge_cells('A1:D1')
            ws_sum['A1'] = f"Resumen — {title} {month} {year}"
            ws_sum['A1'].font = self.font_title

            r = 3
            for key, val in extra_data.items():
                ws_sum.cell(row=r, column=1, value=str(key)).font = Font(name='Calibri', size=10, bold=True)
                if isinstance(val, dict):
                    for k2, v2 in val.items():
                        ws_sum.cell(row=r, column=2, value=str(k2)).font = self.font_data
                        ws_sum.cell(row=r, column=3, value=str(v2)[:100]).font = self.font_data
                        r += 1
                elif isinstance(val, (int, float, str)):
                    ws_sum.cell(row=r, column=2, value=str(val)[:100]).font = self.font_data
                    r += 1
                r += 1

        # ── Hoja 3: Pivotes (solo PQRS) ──
        if tab_type == "pqrs" and extra_data.get("pivots"):
            ws_pivot = wb.create_sheet("TABLAS DINAMICAS")
            ws_pivot.merge_cells('A1:Z1')
            ws_pivot['A1'] = f"Tablas Dinamicas PQRS — {month} {year}"
            ws_pivot['A1'].font = self.font_title

            pr = 3
            for pivot_name, pivot_df in extra_data["pivots"].items():
                if pivot_df is None or pivot_df.empty:
                    continue
                ws_pivot.merge_cells(start_row=pr, start_column=1,
                                     end_row=pr, end_column=min(len(pivot_df.columns), 10))
                ws_pivot.cell(row=pr, column=1, value=pivot_name).font = Font(name='Calibri', size=10, bold=True, color='2E75B6')
                pr += 1
                for col_idx, col_name in enumerate(pivot_df.columns, 1):
                    cell = ws_pivot.cell(row=pr, column=col_idx, value=str(col_name)[:60])
                    cell.font = self.font_hdr
                    cell.fill = self.fill_hdr
                pr += 1
                for _, row_data in pivot_df.iterrows():
                    for col_idx, val in enumerate(row_data, 1):
                        cell = ws_pivot.cell(row=pr, column=col_idx)
                        if pd.notna(val):
                            cell.value = str(val)[:255]
                        cell.font = self.font_data
                    pr += 1
                pr += 1

        # Guardar
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb.save(output_path)
